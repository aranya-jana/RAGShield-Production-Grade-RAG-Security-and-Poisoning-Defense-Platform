"""Secure document upload and ingestion pipeline for RAGShield.

The service performs security checks before a document is added to Chroma.
It intentionally keeps uploaded content out of the persistent document registry
and audit metadata; only security-safe metadata is retained.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from langchain_core.documents import Document

try:
    from src.document_provenance import DocumentProvenanceManager
    from src.document_injection_detector import DocumentInjectionDetector
    from src.poison_detector import PoisonDetector
    from src.pii_detector import PIIDetector
    from src.risk_engine import RiskTrustEngine
except ImportError:
    from document_provenance import DocumentProvenanceManager
    from document_injection_detector import DocumentInjectionDetector
    from poison_detector import PoisonDetector
    from pii_detector import PIIDetector
    from risk_engine import RiskTrustEngine


MAX_TEXT_CHARS = 1_000_000
ALLOWED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}
HIGH_RISK_DLP_CATEGORIES = {
    "credit_card",
    "ssn",
    "aws_access_key",
    "github_token",
    "private_key",
    "jwt",
    "generic_secret",
}


@dataclass
class DocumentRecord:
    document_id: str
    source: str
    extension: str
    size_bytes: int
    status: str
    uploaded_at: str
    content_sha256: str
    metadata_sha256: str
    provenance_version: str
    poison_score: float
    poison_detected: bool
    contradiction_score: float
    contradiction_detected: bool
    injection_score: float
    injection_detected: bool
    dlp_score: float
    dlp_detected: bool
    risk_score: float
    trust_score: float
    classification: str
    detectors: List[str]
    reasons: List[str]


class DocumentRegistry:
    """Small JSON registry containing only non-sensitive document metadata."""

    def __init__(self, path: str = "./data/document_registry.json") -> None:
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._records: Dict[str, DocumentRecord] = {}
        self._load()

    def _load(self) -> None:
        if not self.path.exists():
            return
        try:
            raw = json.loads(self.path.read_text(encoding="utf-8"))
            if isinstance(raw, list):
                for item in raw:
                    if isinstance(item, dict) and item.get("document_id"):
                        self._records[str(item["document_id"])] = DocumentRecord(**item)
        except (OSError, ValueError, TypeError):
            self._records = {}

    def _save(self) -> None:
        payload = [asdict(record) for record in self._records.values()]
        temporary = self.path.with_suffix(".tmp")
        temporary.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        temporary.replace(self.path)

    def upsert(self, record: DocumentRecord) -> None:
        self._records[record.document_id] = record
        self._save()

    def get(self, document_id: str) -> Optional[DocumentRecord]:
        return self._records.get(document_id)

    def list(self) -> List[DocumentRecord]:
        return list(reversed(list(self._records.values())))

    def clear(self) -> None:
        self._records = {}
        try:
            self.path.unlink()
        except FileNotFoundError:
            pass


def sanitize_filename(filename: str) -> str:
    name = Path(filename or "document.txt").name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    if not name:
        name = "document.txt"
    return name[:180]


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def extract_document_text(filename: str, data: bytes) -> Tuple[str, str]:
    """Extract text without executing or rendering uploaded content."""
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported document type. Use TXT, MD, CSV, JSON, PDF, or DOCX.")

    if suffix in {".txt", ".md"}:
        text = _decode_text(data)
    elif suffix == ".csv":
        rows = csv.reader(io.StringIO(_decode_text(data)))
        text = "\n".join(" | ".join(row) for row in rows)
    elif suffix == ".json":
        parsed = json.loads(_decode_text(data))
        text = json.dumps(parsed, indent=2, ensure_ascii=False)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF upload support requires the 'pypdf' package.") from exc
        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    else:  # .docx
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("DOCX upload support requires the 'python-docx' package.") from exc
        doc = DocxDocument(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                table_rows.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(paragraphs + table_rows)

    text = text.replace("\x00", " ").strip()
    if not text:
        raise ValueError("The uploaded document contains no extractable text.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("The extracted document text is too large to process.")
    return text, suffix


def _result_bool(result: Any, name: str) -> bool:
    return bool(getattr(result, name, False))


def _result_score(result: Any) -> float:
    return float(getattr(result, "score", 0.0) or 0.0)


def _result_reasons(result: Any) -> List[str]:
    return list(getattr(result, "reasons", []) or [])


class SecureDocumentIngestionService:
    """Scan, quarantine, or index uploaded documents."""

    def __init__(self, registry: Optional[DocumentRegistry] = None) -> None:
        self.registry = registry or DocumentRegistry()
        self.provenance = DocumentProvenanceManager()
        self.injection_detector = DocumentInjectionDetector()
        self.poison_detector = PoisonDetector()
        self.pii_detector = PIIDetector()
        self.risk_engine = RiskTrustEngine()

    def _metadata_risk(self, document: Document) -> float:
        metadata = dict(document.metadata or {})
        suspicious_keys = {
            "instruction", "instructions", "system_prompt",
            "system_instruction", "prompt", "override", "role",
            "command", "payload",
        }
        suspicious_markers = (
            "ignore previous instructions",
            "ignore all previous instructions",
            "system instruction",
            "critical system instruction",
        )
        for key, value in metadata.items():
            normalized = str(key).strip().lower().replace("-", "_")
            value_text = str(value).strip().lower()
            if normalized in suspicious_keys or any(marker in value_text for marker in suspicious_markers):
                return 1.0
        return 0.0

    def _dlp(self, document: Document) -> Any:
        if hasattr(document, "metadata"):
            text_parts = [document.page_content]
            for key, value in document.metadata.items():
                if not str(key).startswith("ragshield_"):
                    text_parts.append(f"{key}: {value}")
            return self.pii_detector.analyze("\n".join(text_parts))
        return self.pii_detector.analyze(document.page_content)

    @staticmethod
    def _safe_dlp(dlp: Any) -> Dict[str, Any]:
        findings = []
        for finding in getattr(dlp, "findings", []) or []:
            findings.append({
                "category": str(getattr(finding, "category", "unknown")),
                "severity": str(getattr(finding, "severity", "unknown")),
                "reason": str(getattr(finding, "reason", "Sensitive data detected.")),
            })
        return {
            "score": float(getattr(dlp, "score", 0.0) or 0.0),
            "has_pii": bool(getattr(dlp, "has_pii", False)),
            "finding_count": len(findings),
            "categories": sorted({f["category"] for f in findings}),
            "findings": findings,
        }

    def scan_and_ingest(self, *, rag: Any, filename: str, data: bytes, username: str) -> Dict[str, Any]:
        safe_name = sanitize_filename(filename)
        text, suffix = extract_document_text(safe_name, data)
        now = datetime.now(timezone.utc).isoformat()

        document = Document(
            page_content=text,
            metadata={
                "source": safe_name,
                "document_type": "uploaded",
                "ragshield_uploaded_by": username,
                "ragshield_uploaded_at": now,
            },
        )
        document = self.provenance.enrich_document(document)

        injection = self.injection_detector.analyze(text=text, metadata=dict(document.metadata))
        poison = self.poison_detector.analyze(text)
        trusted = [
            d for d in getattr(rag, "trusted_documents", [])
            if getattr(d, "page_content", "") != text
        ]
        contradiction = rag.contradiction_detector.analyze(
            candidate_text=text,
            trusted_documents=trusted,
        )
        dlp = self._dlp(document)
        dlp_categories = {
            str(getattr(f, "category", "unknown"))
            for f in getattr(dlp, "findings", []) or []
        }
        high_risk_dlp = bool(dlp_categories & HIGH_RISK_DLP_CATEGORIES)
        metadata_score = self._metadata_risk(document)

        risk = self.risk_engine.assess(
            injection_score=_result_score(injection),
            poisoning_score=_result_score(poison),
            contradiction_score=_result_score(contradiction),
            metadata_score=metadata_score,
        )

        detectors: List[str] = []
        reasons: List[str] = []
        if _result_bool(injection, "is_injected"):
            detectors.append("DocumentInjectionDetector")
            reasons.extend(_result_reasons(injection))
        if _result_bool(poison, "is_poisoned"):
            detectors.append("PoisonDetector")
            reasons.extend(_result_reasons(poison))
        if _result_bool(contradiction, "is_contradictory"):
            detectors.append("ContradictionDetector")
            reasons.extend(_result_reasons(contradiction))
        if high_risk_dlp:
            detectors.append("PIIDetector")
            reasons.append("High-risk sensitive data detected during document ingestion.")
        if metadata_score >= 1.0:
            detectors.append("MetadataSecurity")
            reasons.append("Instruction-like document metadata detected.")

        reasons = list(dict.fromkeys(reasons))
        blocked = bool(detectors)
        status = "QUARANTINED" if blocked else "INDEXED"

        verification = self.provenance.verify_document(document)

        verification_ok = bool(
            getattr(
                verification,
                "integrity_verified",
                getattr(
                    verification,
                    "is_valid",
                    getattr(
                        verification,
                        "valid",
                        False,
                    ),
                ),
            )
        )

        if not verification_ok:
            blocked = True
            status = "QUARANTINED"
            if "DocumentProvenanceIntegrity" not in detectors:
                detectors.append("DocumentProvenanceIntegrity")
            reasons.append("Document provenance integrity verification failed.")

        record = DocumentRecord(
            document_id=str(document.metadata.get("ragshield_document_id")),
            source=safe_name,
            extension=suffix,
            size_bytes=len(data),
            status=status,
            uploaded_at=now,
            content_sha256=str(document.metadata.get("ragshield_content_sha256", "")),
            metadata_sha256=str(document.metadata.get("ragshield_metadata_sha256", "")),
            provenance_version=str(document.metadata.get("ragshield_provenance_version", "1.0")),
            poison_score=_result_score(poison),
            poison_detected=_result_bool(poison, "is_poisoned"),
            contradiction_score=_result_score(contradiction),
            contradiction_detected=_result_bool(contradiction, "is_contradictory"),
            injection_score=_result_score(injection),
            injection_detected=_result_bool(injection, "is_injected"),
            dlp_score=float(getattr(dlp, "score", 0.0) or 0.0),
            dlp_detected=bool(getattr(dlp, "has_pii", False)),
            risk_score=float(getattr(risk, "risk_score", 0.0)),
            trust_score=float(getattr(risk, "trust_score", 0.0)),
            classification="BLOCKED" if blocked else str(getattr(risk, "classification", "TRUSTED")),
            detectors=list(dict.fromkeys(detectors)),
            reasons=list(dict.fromkeys(reasons)),
        )

        if blocked:
            self.registry.upsert(record)
            return self._response(record, text, dlp, blocked=True)

        # Critical invariant: indexing occurs only after every ingestion check passes.
        rag.vectorstore.add_documents([document])
        if hasattr(rag, "trusted_documents"):
            rag.trusted_documents.append(document)
        if hasattr(rag, "refresh_chain"):
            rag.refresh_chain()
        self.registry.upsert(record)
        return self._response(record, text, dlp, blocked=False)

    def _response(self, record: DocumentRecord, text: str, dlp: Any, *, blocked: bool) -> Dict[str, Any]:
        return {
            "document_id": record.document_id,
            "source": record.source,
            "status": record.status,
            "indexed": not blocked,
            "quarantined": blocked,
            "size_bytes": record.size_bytes,
            "content_sha256": record.content_sha256,
            "provenance_version": record.provenance_version,
            "poison_score": record.poison_score,
            "poison_detected": record.poison_detected,
            "contradiction_score": record.contradiction_score,
            "contradiction_detected": record.contradiction_detected,
            "injection_score": record.injection_score,
            "injection_detected": record.injection_detected,
            "dlp": self._safe_dlp(dlp),
            "risk_score": record.risk_score,
            "trust_score": record.trust_score,
            "classification": record.classification,
            "detectors": record.detectors,
            "reasons": record.reasons,
            "content_preview": text[:500],
        }

    def list_documents(self) -> List[Dict[str, Any]]:
        return [asdict(record) for record in self.registry.list()]

    def clear(self) -> None:
        self.registry.clear()
